#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Génère des fichiers .docx prêts à envoyer à partir des CV et lettres en Markdown.

Usage :
    python3 generer_docx.py

Dépendance : python-docx  (pip install python-docx)
"""

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("python-docx manquant : pip install python-docx")

RACINE = Path(__file__).resolve().parent.parent
BLEU = RGBColor(0x1F, 0x3B, 0x57)
GRIS = RGBColor(0x55, 0x55, 0x55)


# --------------------------------------------------------------------------- #
# Utilitaires de mise en forme
# --------------------------------------------------------------------------- #
def bordure_bas(paragraphe, taille=6, couleur="1F3B57"):
    pPr = paragraphe._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bas = OxmlElement("w:bottom")
    bas.set(qn("w:val"), "single")
    bas.set(qn("w:sz"), str(taille))
    bas.set(qn("w:space"), "2")
    bas.set(qn("w:color"), couleur)
    pbdr.append(bas)
    pPr.append(pbdr)


def ajouter_run(paragraphe, texte, gras=False, italique=False, taille=None, couleur=None):
    """Ajoute un run en interprétant **gras** et *italique* dans le texte."""
    morceaux = re.split(r"(\*\*.+?\*\*|\*.+?\*)", texte)
    for morceau in morceaux:
        if not morceau:
            continue
        g = morceau.startswith("**") and morceau.endswith("**")
        i = (not g) and morceau.startswith("*") and morceau.endswith("*") and len(morceau) > 2
        run = paragraphe.add_run(morceau[2:-2] if g else (morceau[1:-1] if i else morceau))
        run.bold = gras or g
        run.italic = italique or i
        if taille:
            run.font.size = Pt(taille)
        if couleur:
            run.font.color.rgb = couleur


def configurer(doc):
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.4)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.05


# --------------------------------------------------------------------------- #
# Conversion Markdown -> docx
# --------------------------------------------------------------------------- #
def convertir(md_path: Path, docx_path: Path, compact=False):
    lignes = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configurer(doc)
    i = 0

    while i < len(lignes):
        ligne = lignes[i].rstrip()
        strip = ligne.strip()

        # Ligne vide
        if not strip:
            i += 1
            continue

        # Séparateur
        if set(strip) <= {"-", "—"} and len(strip) >= 3:
            p = doc.add_paragraph()
            bordure_bas(p, taille=4, couleur="BBBBBB")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Tableau
        if strip.startswith("|"):
            bloc = []
            while i < len(lignes) and lignes[i].strip().startswith("|"):
                bloc.append(lignes[i].strip())
                i += 1
            rendu = [
                [c.strip() for c in b.strip("|").split("|")]
                for b in bloc
                if not re.fullmatch(r"\|[\s:\-|]+\|", b)
            ]
            if rendu:
                t = doc.add_table(rows=len(rendu), cols=len(rendu[0]))
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r, cellules in enumerate(rendu):
                    for c, valeur in enumerate(cellules[: len(rendu[0])]):
                        cel = t.cell(r, c)
                        cel.text = ""
                        para = cel.paragraphs[0]
                        ajouter_run(para, valeur, gras=(r == 0), taille=8.5)
                doc.add_paragraph()
            continue

        # Titre 1 = nom
        if strip.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(strip[2:].strip())
            run.bold = True
            run.font.size = Pt(19)
            run.font.color.rgb = BLEU
            p.paragraph_format.space_after = Pt(0)
            i += 1
            continue

        # Titre 2 = section
        if strip.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(strip[3:].strip().upper())
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = BLEU
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(3)
            bordure_bas(p, taille=6)
            i += 1
            continue

        # Titre 3
        if strip.startswith("### "):
            p = doc.add_paragraph()
            ajouter_run(p, strip[4:].strip(), gras=True, taille=10.5)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(1)
            i += 1
            continue

        # Citation = note interne (ignorée dans le document final)
        if strip.startswith(">"):
            i += 1
            continue

        # Liste à puces
        if strip.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.space_after = Pt(1)
            ajouter_run(p, strip[2:].strip())
            i += 1
            continue

        # Paragraphe normal
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ajouter_run(p, strip)
        i += 1

    doc.save(docx_path)
    return docx_path


CIBLES = [
    (
        "02_Dossiers/01_PAM_Communication_Assistant_Cotonou/CV_PAM_Communication_Assistant.md",
        "02_Dossiers/01_PAM_Communication_Assistant_Cotonou/CV_GNACADJA_Marie_Edith_Communication_Assistant_PAM.docx",
    ),
    (
        "02_Dossiers/01_PAM_Communication_Assistant_Cotonou/Lettre_Motivation_PAM_Communication_Assistant.md",
        "02_Dossiers/01_PAM_Communication_Assistant_Cotonou/Lettre_Motivation_GNACADJA_PAM_Communication_Assistant.docx",
    ),
    (
        "02_Dossiers/02_FAO_Stage_Volontariat_Afrique/CV_FAO_Nutrition_Securite_Alimentaire.md",
        "02_Dossiers/02_FAO_Stage_Volontariat_Afrique/CV_GNACADJA_Marie_Edith_FAO_Nutrition.docx",
    ),
    (
        "02_Dossiers/02_FAO_Stage_Volontariat_Afrique/Lettre_Motivation_FAO_Stage_Afrique.md",
        "02_Dossiers/02_FAO_Stage_Volontariat_Afrique/Lettre_Motivation_GNACADJA_FAO.docx",
    ),
]

if __name__ == "__main__":
    for source, cible in CIBLES:
        src = RACINE / source
        dst = RACINE / cible
        if not src.exists():
            print(f"absent : {src}")
            continue
        convertir(src, dst)
        print(f"généré : {dst.relative_to(RACINE)}")
